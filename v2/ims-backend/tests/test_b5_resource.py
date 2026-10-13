"""B5 公司资源：公司、实名人、手机卡、证件。"""
from __future__ import annotations

import base64
import io
from datetime import timedelta

from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import func, select

from app.core.clock import today
from app.core.context import system_context, use_context
from app.db.session import session_scope
from app.domains.cert.model import CertExpireTask, CertViewLog
from app.domains.cert.service import is_locked
from app.domains.master.model import Cert, Realname
from app.domains.message.model import Message, NotifyRecord
from app.domains.resource.common import id_check_char
from app.domains.system.model import OperateLog, SensitiveViewLog
from tests.conftest import auth
from tests.stddata import U_ADMIN, U_DEPT, U_RESIGNED

API = "/admin-api/ims"
_CC = "0123456789ABCDEFGHJKLMNPQRTUWXY"
_CW = [1, 3, 9, 27, 19, 26, 16, 17, 20, 29, 25, 13, 8, 24, 10, 30, 28]


def credit(body: str) -> str:
    text = body[:17].upper()
    total = sum(_CC.index(ch) * weight for ch, weight in zip(text, _CW))
    return text + _CC[(31 - total % 31) % 31]


def id_card(body: str) -> str:
    return body[:17] + id_check_char(body[:17])


def png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (2, 2), (10, 20, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


def day(offset: int) -> str:
    return (today() + timedelta(days=offset)).isoformat()


def walk(nodes):
    for node in nodes or []:
        yield node
        yield from walk(node.get("children") or [])


def sheet_text(client, headers, payload: dict) -> str:
    assert payload["status"] == "SUCCESS"
    downloaded = client.get(f"{API}/export-task/download", headers=headers, params={"id": payload["id"]}).json()
    assert downloaded["code"] == 0, downloaded
    book = load_workbook(io.BytesIO(base64.b64decode(downloaded["data"]["content"])))
    cells = []
    for row in book.active.iter_rows():
        cells.extend("" if cell.value is None else str(cell.value) for cell in row)
    return "\n".join(cells)


def test_company_realname_sim_and_selectors(client, monkeypatch):
    admin = auth(client, "u_admin")
    empty = client.get(f"{API}/master/selector/company", headers=admin).json()
    assert empty["code"] == 1221

    created = client.post(
        f"{API}/corp/resource/company",
        headers=admin,
        json={
            "name": "星河传媒",
            "creditCode": credit("91310000MA1FL2XW3"),
            "legalPerson": "周宁",
            "address": "上海",
            "wxOfficialCapacity": 5,
            "clientToken": "co-1",
        },
    ).json()
    assert created["code"] == 0, created
    company_id = created["data"]["id"]
    assert isinstance(company_id, str)
    assert created["data"]["companyNo"].startswith("CO")
    again = client.post(
        f"{API}/corp/resource/company",
        headers=admin,
        json={"name": "另一家", "creditCode": credit("91310000MA1FL2XW4"), "clientToken": "co-1"},
    ).json()
    assert again["data"]["id"] == company_id

    page = client.get(f"{API}/corp/resource/company/page", headers=admin).json()
    assert page["data"]["pageSize"] == 20
    assert page["data"]["total"] == 1
    detail = client.get(f"{API}/corp/resource/company/{company_id}", headers=admin).json()
    assert detail["data"]["capacityText"] == "0 / 5"
    renamed = client.put(f"{API}/corp/resource/company/{company_id}", headers=admin, json={"name": "星河文化"}).json()
    assert renamed["data"]["name"] == "星河文化"
    stopped = client.put(f"{API}/corp/resource/company/{company_id}/status", headers=admin, json={"status": "DISABLED"}).json()
    assert stopped["data"]["status"] == "DISABLED"
    client.put(f"{API}/corp/resource/company/{company_id}/status", headers=admin, json={"status": "ENABLED"})

    monkeypatch.setattr("app.domains.resource.company.wechat_official_used", lambda _ident: 3)
    tight = client.put(
        f"{API}/corp/resource/company/{company_id}/capacity",
        headers=admin,
        json={"capacity": 2, "reason": "缩容"},
    ).json()
    assert tight["code"] == 1222
    assert "3" in tight["msg"] and "5" in tight["msg"]
    opened = client.put(
        f"{API}/corp/resource/company/{company_id}/capacity",
        headers=admin,
        json={"capacity": 8, "reason": "扩容"},
    ).json()
    assert opened["code"] == 0
    assert opened["data"]["capacityLogs"][0]["title"] == "5 → 8"

    exported = client.get(f"{API}/corp/resource/company/export", headers=admin).json()
    assert "星河文化" in sheet_text(client, admin, exported["data"])
    denied = client.get(f"{API}/corp/resource/company/page", headers=auth(client, "u_self")).json()
    assert denied["code"] == 1008

    person = client.post(
        f"{API}/corp/resource/realname",
        headers=admin,
        json={"name": "林夏", "idType": "ID_CARD", "idNo": id_card("11010119900307851"), "mobile": "13800138000", "userId": str(U_DEPT)},
    ).json()
    assert person["code"] == 0, person
    person_id = person["data"]["id"]
    assert "****" in person["data"]["idNoMasked"]
    assert person["data"]["idNoMasked"] != id_card("11010119900307851")
    duplicate = client.post(
        f"{API}/corp/resource/realname",
        headers=admin,
        json={"name": "另一个", "idType": "ID_CARD", "idNo": id_card("11010119900307851"), "userId": str(U_ADMIN)},
    ).json()
    assert duplicate["code"] == 1001
    resigned = client.post(
        f"{API}/corp/resource/realname",
        headers=admin,
        json={"name": "离职关联", "idType": "ID_CARD", "idNo": id_card("11010119900101123"), "userId": str(U_RESIGNED)},
    ).json()
    assert resigned["code"] == 1501
    plain = client.get(f"{API}/corp/resource/realname/{person_id}/plain", headers=admin).json()
    assert plain["data"]["idNo"] == id_card("11010119900307851")
    masked_book = sheet_text(client, admin, client.get(f"{API}/corp/resource/realname/export", headers=admin).json()["data"])
    assert id_card("11010119900307851") not in masked_book
    assert "****" in masked_book

    sim = client.post(
        f"{API}/corp/resource/sim-card",
        headers=admin,
        json={
            "phone": "13900001111",
            "operator": "CMCC",
            "iccid": "8986001234567890123",
            "realnameId": person_id,
            "monthlyRent": "12.50",
            "paymentCycle": "MONTH",
            "activatedAt": "2026-01-01",
            "nextPaymentDate": day(-1),
        },
    ).json()
    assert sim["code"] == 0, sim
    sim_id = sim["data"]["id"]
    assert sim["data"]["phoneMasked"].endswith("1111")
    assert sim["data"]["monthlyRent"] == "12.50"
    assert sim["data"]["paymentTone"] == "danger"
    assert sim["data"]["phoneName"] == "—"
    revealed = client.get(f"{API}/corp/resource/sim-card/{sim_id}/plain", headers=admin).json()
    assert revealed["data"]["phone"] == "13900001111"
    listed = client.get(f"{API}/corp/resource/sim-card/page", headers=admin, params={"phone": "13900001111"}).json()
    assert listed["data"]["total"] == 1
    blocked = client.delete(f"{API}/corp/resource/realname/{person_id}", headers=admin).json()
    assert blocked["code"] == 1502
    removed = client.delete(f"{API}/corp/resource/sim-card/{sim_id}", headers=admin).json()
    assert removed["code"] == 0
    client.delete(f"{API}/corp/resource/realname/{person_id}", headers=admin)
    client.delete(f"{API}/corp/resource/company/{company_id}", headers=admin)

    picked = client.get(f"{API}/master/selector/company", headers=admin, params={"keyword": "不存在的公司"}).json()
    assert picked["code"] == 0
    assert picked["data"]["list"] == []


def test_certificate_review_level_scan_and_risk(client):
    admin = auth(client, "u_admin")
    me = client.get(f"{API}/auth/me", headers=admin).json()
    names = {node["name"] for node in walk(me["data"]["menus"])}
    assert {"公司管理", "实名人管理", "手机卡管理", "证件管理", "账号管理", "公众号"} <= names
    assert "ims:cert:archive:upload" in me["data"]["permCodes"]

    person = client.post(
        f"{API}/corp/resource/realname",
        headers=admin,
        json={"name": "陈舟", "idType": "ID_CARD", "idNo": id_card("11010119920202333"), "userId": str(U_DEPT)},
    ).json()
    assert person["code"] == 0, person
    person_id = person["data"]["id"]
    outsider = client.post(
        f"{API}/corp/resource/realname",
        headers=admin,
        json={"name": "外部持有", "idType": "PASSPORT", "idNo": "E12345678"},
    ).json()
    assert outsider["code"] == 0, outsider

    uploaded = client.post(
        f"{API}/file/upload",
        headers={"Authorization": admin["Authorization"]},
        data={"scene": "CERT"},
        files={"file": ("id.png", png(), "image/png")},
    ).json()
    assert uploaded["code"] == 0, uploaded
    assert uploaded["data"]["ocr"]["ok"] is True
    file_key = uploaded["data"]["fileKey"]

    pending = client.post(
        f"{API}/cert/archive/upload",
        headers=admin,
        json={
            "realnameId": person_id,
            "certType": "ID_CARD",
            "certNoPlain": id_card("11010119920202333"),
            "issueDate": "2020-01-01",
            "expireDate": day(10),
            "fileKey": file_key,
            "clientToken": "cert-1",
        },
    ).json()
    assert pending["code"] == 0, pending
    cert_id = pending["data"]["id"]
    assert isinstance(cert_id, str)
    replay = client.post(
        f"{API}/cert/archive/upload",
        headers=admin,
        json={"clientToken": "cert-1", "realnameId": person_id},
    ).json()
    assert replay["data"]["id"] == cert_id

    blocked_view = client.get(f"{API}/cert/security/view/{cert_id}", headers=admin).json()
    assert blocked_view["code"] == 1033
    self_review = client.put(f"{API}/cert/archive/{cert_id}/review", headers=admin, json={"action": "APPROVE"}).json()
    assert self_review["code"] == 1008

    with session_scope(commit=True) as session:
        with use_context(system_context(1)):
            session.get(Cert, int(cert_id)).uploaded_by = U_DEPT
    approved = client.put(f"{API}/cert/archive/{cert_id}/review", headers=admin, json={"action": "APPROVE"}).json()
    assert approved["code"] == 0, approved
    assert approved["data"]["status"] == "EXPIRING"

    duplicate = client.post(
        f"{API}/cert/archive/upload",
        headers=admin,
        json={
            "realnameId": person_id,
            "certType": "ID_CARD",
            "certNoPlain": id_card("11010119920202999"),
            "issueDate": "2024-01-01",
            "expireDate": "2034-01-01",
            "fileKey": file_key,
        },
    ).json()
    assert duplicate["code"] == 1032

    level = client.get(f"{API}/cert/security/level-config", headers=admin).json()
    assert level["code"] == 0
    stale = client.put(
        f"{API}/cert/security/level-config",
        headers=admin,
        json={"version": 0, "l2RoleIds": ["101"], "opacity": "0.12", "position": "BOTTOM_RIGHT"},
    ).json()
    assert stale["code"] == 1001
    saved = client.put(
        f"{API}/cert/security/level-config",
        headers=admin,
        json={"version": level["data"]["version"], "l2RoleIds": ["101"], "l3UserIds": [], "opacity": "0.12", "position": "BOTTOM_RIGHT"},
    ).json()
    assert saved["code"] == 0, saved
    assert saved["data"]["hint"] == "下次查看即按新级别"

    viewed = client.get(f"{API}/cert/security/view/{cert_id}", headers=admin).json()
    assert viewed["code"] == 0, viewed
    assert viewed["data"]["viewLevel"] == "L2"
    assert viewed["data"]["watermarkText"].startswith("管理员 u_admin ")
    streamed = client.get(viewed["data"]["fileUrl"], headers=admin)
    assert streamed.status_code == 200
    assert streamed.headers["content-type"].startswith("image/")
    assert streamed.content != png()
    duration = client.get(
        f"{API}/cert/security/view/{cert_id}",
        headers=admin,
        params={"logId": viewed["data"]["logId"], "durationSec": 8},
    ).json()
    assert duration["data"]["durationSec"] == 8

    for _ in range(10):
        opened = client.get(f"{API}/cert/security/view/{cert_id}", headers=admin).json()
        assert opened["code"] == 0
    report = client.get(f"{API}/cert/security/risk-report", headers=admin).json()
    assert report["data"]["list"]
    exported = client.get(f"{API}/cert/security/risk-report", headers=admin, params={"export": "1"}).json()
    assert exported["data"]["fileKey"].endswith("证件异常访问.xlsx")
    with session_scope() as session:
        with use_context(system_context(1)):
            views = session.scalar(select(func.count()).select_from(CertViewLog))
            assert int(views or 0) == 11
            assert session.scalar(select(OperateLog.id).where(OperateLog.action == "异常报告", OperateLog.method == "GET")) is not None

    tasks = client.get(f"{API}/cert/expire/list", headers=admin).json()
    assert tasks["data"]["total"] >= 1
    task_id = tasks["data"]["list"][0]["id"]
    with session_scope() as session:
        with use_context(system_context(1)):
            before = session.scalar(select(func.count()).select_from(NotifyRecord).where(NotifyRecord.biz_key == f"cert:{cert_id}:YELLOW"))
    scanned = client.post(f"{API}/cert/expire/scan", headers=admin).json()
    assert scanned["code"] == 0, scanned
    client.post(f"{API}/cert/expire/scan", headers=admin)
    with session_scope() as session:
        with use_context(system_context(1)):
            after = session.scalar(select(func.count()).select_from(NotifyRecord).where(NotifyRecord.biz_key == f"cert:{cert_id}:YELLOW"))
    assert int(after or 0) == int(before or 0)
    stats = client.get(f"{API}/cert/expire/stats", headers=admin).json()
    assert stats["data"]["yellow"] >= 1
    missing = client.put(f"{API}/cert/expire/999999/remind", headers=admin, json={"channel": "SITE"}).json()
    assert missing["code"] == 1039
    reminded = client.put(f"{API}/cert/expire/{task_id}/remind", headers=admin, json={"channel": "SITE"}).json()
    assert reminded["code"] == 0, reminded
    assert reminded["data"]["remindCount"] == 1
    with session_scope() as session:
        with use_context(system_context(1)):
            site = session.scalar(select(NotifyRecord.id).where(NotifyRecord.biz_key == f"cert-remind:{task_id}:1", NotifyRecord.channel == "SITE"))
            ding = session.scalar(select(NotifyRecord.id).where(NotifyRecord.biz_key == f"cert-remind:{task_id}:1", NotifyRecord.channel == "DINGTALK"))
            route = session.scalar(select(Message.route).where(Message.route.contains(f"certId={cert_id}")))
    assert site is not None and ding is None
    assert route and "tab=expire" in route

    renewed_file = client.post(
        f"{API}/file/upload",
        headers={"Authorization": admin["Authorization"]},
        data={"scene": "CERT"},
        files={"file": ("new.png", png(), "image/png")},
    ).json()["data"]["fileKey"]
    newer = client.post(
        f"{API}/cert/archive/upload",
        headers=admin,
        json={
            "realnameId": person_id,
            "certType": "ID_CARD",
            "certNoPlain": id_card("11010119920202999"),
            "issueDate": "2024-01-01",
            "expireDate": "2034-01-01",
            "fileKey": renewed_file,
            "renewOfCertId": cert_id,
        },
    ).json()
    assert newer["code"] == 0, newer
    new_id = newer["data"]["id"]
    bound = client.put(f"{API}/cert/expire/{task_id}/renew", headers=admin, json={"newCertId": new_id}).json()
    assert bound["code"] == 0, bound
    assert "新证需审核生效后旧证自动归档" in bound["data"]["hint"]
    with session_scope(commit=True) as session:
        with use_context(system_context(1)):
            session.get(Cert, int(new_id)).uploaded_by = U_DEPT
    finished = client.put(f"{API}/cert/archive/{new_id}/review", headers=admin, json={"action": "APPROVE"}).json()
    assert finished["data"]["status"] == "EFFECTIVE"
    old = client.get(f"{API}/cert/archive/{cert_id}", headers=admin).json()
    assert old["code"] == 0, old
    assert old["data"]["status"] == "RECYCLED"
    with session_scope() as session:
        with use_context(system_context(1)):
            task = session.get(CertExpireTask, int(task_id))
            assert task.status == "RENEW_RESOLVED"

    metrics = client.get(f"{API}/cert/archive/digital-metrics", headers=admin).json()
    assert metrics["data"]["denominator"] == 2
    assert metrics["data"]["numerator"] == 1
    assert metrics["data"]["belowStandard"] is True

    locked_file = client.post(
        f"{API}/file/upload",
        headers={"Authorization": admin["Authorization"]},
        data={"scene": "CERT"},
        files={"file": ("old.png", png(), "image/png")},
    ).json()["data"]["fileKey"]
    expired = client.post(
        f"{API}/cert/archive/upload",
        headers=admin,
        json={
            "realnameId": outsider["data"]["id"],
            "certType": "PASSPORT",
            "certNoPlain": "E12345678",
            "issueDate": "2020-01-01",
            "expireDate": day(-1),
            "fileKey": locked_file,
        },
    ).json()
    assert expired["code"] == 0, expired
    with session_scope(commit=True) as session:
        with use_context(system_context(1)):
            session.get(Cert, int(expired["data"]["id"])).uploaded_by = U_DEPT
    locked = client.put(f"{API}/cert/archive/{expired['data']['id']}/review", headers=admin, json={"action": "APPROVE"}).json()
    assert locked["data"]["status"] == "EXPIRED"
    with session_scope() as session:
        with use_context(system_context(1)):
            assert is_locked(int(outsider["data"]["id"])) is True
            assert session.scalar(select(SensitiveViewLog.id).where(SensitiveViewLog.biz_type == "CERT")) is not None

    long_file = client.post(
        f"{API}/file/upload",
        headers={"Authorization": admin["Authorization"]},
        data={"scene": "CERT"},
        files={"file": ("long.png", png(), "image/png")},
    ).json()["data"]["fileKey"]
    lasting = client.post(
        f"{API}/cert/archive/upload",
        headers=admin,
        json={
            "realnameId": outsider["data"]["id"],
            "certType": "HEALTH",
            "certNoPlain": "HEALTH-1",
            "issueDate": "2020-01-01",
            "longTerm": True,
            "fileKey": long_file,
        },
    ).json()
    assert lasting["code"] == 0, lasting
    assert lasting["data"]["remainText"] == "长期"
    with session_scope(commit=True) as session:
        with use_context(system_context(1)):
            session.get(Cert, int(lasting["data"]["id"])).uploaded_by = U_DEPT
    client.put(f"{API}/cert/archive/{lasting['data']['id']}/review", headers=admin, json={"action": "APPROVE"})
    with session_scope() as session:
        with use_context(system_context(1)):
            assert session.scalar(select(CertExpireTask.id).where(CertExpireTask.cert_id == int(lasting["data"]["id"]))) is None

    logs = client.get(f"{API}/cert/security/view-logs", headers=admin).json()
    assert logs["data"]["total"] >= 1
    rejected_holder = client.post(
        f"{API}/corp/resource/realname",
        headers=admin,
        json={"name": "驳回样例", "idType": "ID_CARD", "idNo": id_card("11010119880101666")},
    ).json()
    reject_file = client.post(
        f"{API}/file/upload",
        headers={"Authorization": admin["Authorization"]},
        data={"scene": "CERT"},
        files={"file": ("bad.png", png(), "image/png")},
    ).json()["data"]["fileKey"]
    bad = client.post(
        f"{API}/cert/archive/upload",
        headers=admin,
        json={
            "realnameId": rejected_holder["data"]["id"],
            "certType": "ID_CARD",
            "certNoPlain": id_card("11010119880101666"),
            "issueDate": "2020-01-01",
            "expireDate": "2030-01-01",
            "fileKey": reject_file,
        },
    ).json()
    with session_scope(commit=True) as session:
        with use_context(system_context(1)):
            session.get(Cert, int(bad["data"]["id"])).uploaded_by = U_DEPT
    rejected = client.put(
        f"{API}/cert/archive/{bad['data']['id']}/review",
        headers=admin,
        json={"action": "REJECT", "reason": "影像模糊"},
    ).json()
    assert rejected["code"] == 0, rejected
    assert rejected["data"]["status"] == "DELETED"
    gone = client.get(f"{API}/cert/archive/{bad['data']['id']}", headers=admin).json()
    assert gone["code"] != 0
    with session_scope() as session:
        with use_context(system_context(1)):
            note = session.scalar(select(Message).where(Message.receiver_id == U_DEPT, Message.route.contains("tab=archive")))
    assert note is not None
